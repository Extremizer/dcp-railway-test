#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""Stock file inspection/import helpers for Extremizer Pro warehouses."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
import xlrd

HEADER_SCAN_ROWS = 20
PREVIEW_ROWS = 8


def _env_int(name: str, default: int) -> int:
    try:
        value = int(str(os.getenv(name, default)).strip())
        return value if value > 0 else default
    except (TypeError, ValueError):
        return default


MAX_FILE_BYTES = _env_int(
    "WAREHOUSE_FILE_MAX_BYTES",
    25 * 1024 * 1024,
)
MAX_UNCOMPRESSED_BYTES = _env_int(
    "WAREHOUSE_FILE_MAX_UNCOMPRESSED_BYTES",
    200 * 1024 * 1024,
)
MAX_ROWS = _env_int(
    "WAREHOUSE_FILE_MAX_ROWS",
    100_000,
)
MAX_COLUMNS = _env_int(
    "WAREHOUSE_FILE_MAX_COLUMNS",
    200,
)
MAX_SHEETS = _env_int(
    "WAREHOUSE_FILE_MAX_SHEETS",
    25,
)
MAX_WORKBOOK_ROWS = _env_int(
    "WAREHOUSE_FILE_MAX_WORKBOOK_ROWS",
    250_000,
)
MAX_CELLS = _env_int(
    "WAREHOUSE_FILE_MAX_CELLS",
    1_000_000,
)
MAX_CELL_CHARS = _env_int(
    "WAREHOUSE_FILE_MAX_CELL_CHARS",
    10_000,
)
MAX_ZIP_ENTRIES = _env_int(
    "WAREHOUSE_FILE_MAX_ZIP_ENTRIES",
    2_000,
)
MAX_COMPRESSION_RATIO = _env_int(
    "WAREHOUSE_FILE_MAX_COMPRESSION_RATIO",
    250,
)


class WarehouseFileError(ValueError):
    """Safe, user-facing warehouse file validation error."""


class WarehouseFileLimitError(WarehouseFileError):
    """Warehouse file exceeds a configured safety limit."""


def _mb(value: int) -> str:
    return f"{value / (1024 * 1024):.1f} MB"


def safety_limits() -> dict[str, int]:
    return {
        "max_file_bytes": MAX_FILE_BYTES,
        "max_uncompressed_bytes": MAX_UNCOMPRESSED_BYTES,
        "max_rows": MAX_ROWS,
        "max_columns": MAX_COLUMNS,
        "max_sheets": MAX_SHEETS,
        "max_workbook_rows": MAX_WORKBOOK_ROWS,
        "max_cells": MAX_CELLS,
        "max_cell_chars": MAX_CELL_CHARS,
        "max_zip_entries": MAX_ZIP_ENTRIES,
        "max_compression_ratio": MAX_COMPRESSION_RATIO,
    }


def _check_cell(value: Any) -> None:
    if value is None:
        return
    if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
        raise WarehouseFileLimitError(
            "В файле найдена слишком длинная ячейка: "
            f"более {MAX_CELL_CHARS} символов."
        )


def _check_row_width(row: list[Any], row_number: int | None = None) -> None:
    if len(row) > MAX_COLUMNS:
        place = (
            f" в строке {row_number}"
            if row_number is not None
            else ""
        )
        raise WarehouseFileLimitError(
            f"Слишком много колонок{place}: "
            f"{len(row)}. Лимит: {MAX_COLUMNS}."
        )
    for value in row:
        _check_cell(value)


def preflight_file(path: str | Path) -> dict[str, Any]:
    file_path = Path(path)
    if not file_path.exists() or not file_path.is_file():
        raise WarehouseFileError("Файл не найден.")

    suffix = file_path.suffix.lower()
    if suffix not in {".xlsx", ".xls", ".csv"}:
        raise WarehouseFileError(
            "Поддерживаются только XLSX, XLS и CSV."
        )

    size = file_path.stat().st_size
    if size <= 0:
        raise WarehouseFileError("Файл пустой.")
    if size > MAX_FILE_BYTES:
        raise WarehouseFileLimitError(
            "Файл слишком большой: "
            f"{_mb(size)}. Лимит: {_mb(MAX_FILE_BYTES)}."
        )

    result: dict[str, Any] = {
        "file_size": size,
        "file_format": suffix.lstrip("."),
    }

    if suffix == ".xlsx":
        try:
            with zipfile.ZipFile(file_path) as archive:
                infos = archive.infolist()
                if len(infos) > MAX_ZIP_ENTRIES:
                    raise WarehouseFileLimitError(
                        "XLSX содержит слишком много внутренних файлов: "
                        f"{len(infos)}. Лимит: {MAX_ZIP_ENTRIES}."
                    )

                uncompressed = sum(
                    max(0, int(info.file_size))
                    for info in infos
                )
                compressed = sum(
                    max(0, int(info.compress_size))
                    for info in infos
                )
                result["uncompressed_bytes"] = uncompressed

                if uncompressed > MAX_UNCOMPRESSED_BYTES:
                    raise WarehouseFileLimitError(
                        "Распакованный XLSX слишком большой: "
                        f"{_mb(uncompressed)}. Лимит: "
                        f"{_mb(MAX_UNCOMPRESSED_BYTES)}."
                    )

                if compressed > 0:
                    ratio = uncompressed / compressed
                    result["compression_ratio"] = ratio
                    if ratio > MAX_COMPRESSION_RATIO:
                        raise WarehouseFileLimitError(
                            "XLSX имеет подозрительно высокий коэффициент "
                            f"сжатия: {ratio:.0f}×. Лимит: "
                            f"{MAX_COMPRESSION_RATIO}×."
                        )
        except zipfile.BadZipFile as exc:
            raise WarehouseFileError(
                "Файл XLSX повреждён или имеет неверный формат."
            ) from exc

    return result

OEM_NAMES = {
    "oem", "sku", "partnumber", "partno", "article", "artikl",
    "артикул", "каталожныйномер", "катномер", "номенклатуракод",
    "кодтовара", "код", "партномер", "номердетали",
}
QTY_NAMES = {
    "qty", "quantity", "stock", "available", "balance", "onhand",
    "остаток", "остатки", "количество", "колво", "наличие",
    "свободно", "доступно", "склад", "свободныйостаток",
}
QTY_PREFERRED_NAMES = {
    "свободныйостаток", "freestock", "freequantity",
    "availablequantity", "availableqty",
}
MFG_NAMES = {
    "manufacturer", "brand", "make", "производитель", "бренд", "марка",
}
NAME_NAMES = {
    "name", "description", "product", "partname", "название",
    "наименование", "номенклатура", "товар", "описание",
}
PRICE_NAMES = {
    "price", "unitprice", "saleprice", "retailprice", "цена",
    "стоимость", "ценаруб", "стоимостьруб",
}


def _norm(value: Any) -> str:
    text = str(value or "").strip().lower().replace("ё", "е")
    return re.sub(r"[^a-zа-я0-9]+", "", text)


def header_signature(headers: list[Any]) -> str:
    normalized = [_norm(value) for value in headers]
    payload = json.dumps(
        normalized,
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(payload).hexdigest()


def _clean_row(row: list[Any]) -> list[Any]:
    values = list(row)
    while values and (values[-1] is None or str(values[-1]).strip() == ""):
        values.pop()
    return values


def _score_header(row: list[Any]) -> tuple[int, dict[str, int]]:
    mapping: dict[str, int] = {}
    preferred_quantity_index = None
    for index, value in enumerate(row):
        key = _norm(value)
        if not key:
            continue
        if key in QTY_PREFERRED_NAMES and preferred_quantity_index is None:
            preferred_quantity_index = index
        if key in OEM_NAMES and "oem" not in mapping:
            mapping["oem"] = index
        elif key in QTY_NAMES and "quantity" not in mapping:
            mapping["quantity"] = index
        elif key in MFG_NAMES and "manufacturer" not in mapping:
            mapping["manufacturer"] = index
        elif key in NAME_NAMES and "name" not in mapping:
            mapping["name"] = index
        elif key in PRICE_NAMES and "price" not in mapping:
            mapping["price"] = index
    if preferred_quantity_index is not None:
        mapping["quantity"] = preferred_quantity_index

    score = 0
    if "oem" in mapping:
        score += 5
    if "quantity" in mapping:
        score += 5
    if "manufacturer" in mapping:
        score += 1
    if "name" in mapping:
        score += 1
    return score, mapping


def _detect_header(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    best_index = 0
    best_score = -1
    best_mapping: dict[str, int] = {}
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        score, mapping = _score_header(row)
        if score > best_score:
            best_index = index
            best_score = score
            best_mapping = mapping
    return best_index, best_mapping


def _read_csv(path: Path) -> tuple[str, list[list[Any]]]:
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1251", "utf-8"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("latin1")

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"

    rows: list[list[Any]] = []
    cell_count = 0
    for row_number, raw_row in enumerate(
        csv.reader(text.splitlines(), dialect),
        start=1,
    ):
        if row_number > MAX_ROWS:
            raise WarehouseFileLimitError(
                f"В CSV больше {MAX_ROWS} строк."
            )
        row = _clean_row(raw_row)
        _check_row_width(row, row_number)
        cell_count += len(row)
        if cell_count > MAX_CELLS:
            raise WarehouseFileLimitError(
                f"В CSV слишком много ячеек. Лимит: {MAX_CELLS}."
            )
        rows.append(row)
    return "CSV", rows


def _read_xlsx(
    path: Path,
    sheet_name: str | None = None,
) -> tuple[str, list[list[Any]]]:
    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        if sheet_name and sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]

        declared_rows = int(sheet.max_row or 0)
        declared_columns = int(sheet.max_column or 0)
        if declared_rows > MAX_ROWS:
            raise WarehouseFileLimitError(
                f"Лист «{sheet.title}» заявляет {declared_rows} строк. "
                f"Лимит: {MAX_ROWS}."
            )
        if declared_columns > MAX_COLUMNS:
            raise WarehouseFileLimitError(
                f"Лист «{sheet.title}» заявляет {declared_columns} колонок. "
                f"Лимит: {MAX_COLUMNS}."
            )

        rows: list[list[Any]] = []
        cell_count = 0
        for row_number, raw_row in enumerate(
            sheet.iter_rows(values_only=True),
            start=1,
        ):
            if row_number > MAX_ROWS:
                raise WarehouseFileLimitError(
                    f"На листе «{sheet.title}» больше {MAX_ROWS} строк."
                )
            row = _clean_row(list(raw_row))
            _check_row_width(row, row_number)
            cell_count += len(row)
            if cell_count > MAX_CELLS:
                raise WarehouseFileLimitError(
                    f"На листе «{sheet.title}» слишком много ячеек. "
                    f"Лимит: {MAX_CELLS}."
                )
            rows.append(row)
        return sheet.title, rows
    finally:
        workbook.close()


def _read_xls(
    path: Path,
    sheet_name: str | None = None,
) -> tuple[str, list[list[Any]]]:
    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        if sheet_name and sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
        else:
            sheet = workbook.sheet_by_index(0)

        if sheet.nrows > MAX_ROWS:
            raise WarehouseFileLimitError(
                f"На листе «{sheet.name}» больше {MAX_ROWS} строк."
            )
        if sheet.ncols > MAX_COLUMNS:
            raise WarehouseFileLimitError(
                f"На листе «{sheet.name}» слишком много колонок: "
                f"{sheet.ncols}. Лимит: {MAX_COLUMNS}."
            )
        if int(sheet.nrows) * int(sheet.ncols) > MAX_CELLS:
            raise WarehouseFileLimitError(
                f"На листе «{sheet.name}» слишком много ячеек. "
                f"Лимит: {MAX_CELLS}."
            )

        rows: list[list[Any]] = []
        for row_index in range(sheet.nrows):
            row = _clean_row([
                sheet.cell_value(row_index, c)
                for c in range(sheet.ncols)
            ])
            _check_row_width(row, row_index + 1)
            rows.append(row)
        return sheet.name, rows
    finally:
        try:
            workbook.release_resources()
        except Exception:
            pass


def _sheet_names(path: Path) -> list[str]:
    preflight_file(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return ["CSV"]
    if suffix == ".xlsx":
        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        try:
            names = list(workbook.sheetnames)
            if len(names) > MAX_SHEETS:
                raise WarehouseFileLimitError(
                    f"В XLSX слишком много листов: {len(names)}. "
                    f"Лимит: {MAX_SHEETS}."
                )
            return names
        finally:
            workbook.close()
    if suffix == ".xls":
        workbook = xlrd.open_workbook(path, on_demand=True)
        try:
            names = list(workbook.sheet_names())
            if len(names) > MAX_SHEETS:
                raise WarehouseFileLimitError(
                    f"В XLS слишком много листов: {len(names)}. "
                    f"Лимит: {MAX_SHEETS}."
                )
            return names
        finally:
            try:
                workbook.release_resources()
            except Exception:
                pass
    raise WarehouseFileError("Поддерживаются только XLSX, XLS и CSV.")


def read_table(
    path: str | Path,
    sheet_name: str | None = None,
) -> tuple[str, str, list[list[Any]]]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        sheet, rows = _read_csv(file_path)
        return "csv", sheet, rows
    if suffix == ".xlsx":
        sheet, rows = _read_xlsx(file_path, sheet_name)
        return "xlsx", sheet, rows
    if suffix == ".xls":
        sheet, rows = _read_xls(file_path, sheet_name)
        return "xls", sheet, rows
    raise ValueError("Поддерживаются только XLSX, XLS и CSV.")


def list_sheet_names(path: str | Path) -> list[str]:
    return _sheet_names(Path(path))


def inspect_stock_file(
    path: str | Path,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    file_path = Path(path)
    best: dict[str, Any] | None = None
    available_sheets = _sheet_names(file_path)
    requested_sheets = (
        [sheet_name]
        if sheet_name and sheet_name in available_sheets
        else available_sheets
    )

    workbook_rows_scanned = 0

    for sheet_order, requested_sheet in enumerate(requested_sheets):
        file_format, sheet_name, rows = read_table(
            file_path,
            None if file_path.suffix.lower() == ".csv" else requested_sheet,
        )
        workbook_rows_scanned += len(rows)
        if workbook_rows_scanned > MAX_WORKBOOK_ROWS:
            raise WarehouseFileLimitError(
                "Во всей книге слишком много строк для автоматического "
                f"сканирования: больше {MAX_WORKBOOK_ROWS}."
            )

        if not any(
            any(str(x or "").strip() for x in row)
            for row in rows
        ):
            continue

        header_index, mapping = _detect_header(rows)
        header = [str(x or "").strip() for x in rows[header_index]]
        header_score, _ = _score_header(rows[header_index])
        data_rows = [
            row
            for row in rows[header_index + 1:]
            if any(str(x or "").strip() for x in row)
        ]

        candidate = {
            "file_format": file_format,
            "sheet_name": sheet_name,
            "available_sheets": list(available_sheets),
            "sheet_order": sheet_order,
            "header_row": header_index + 1,
            "headers": header,
            "header_signature": header_signature(header),
            "mapping": mapping,
            "mapping_ready": (
                "oem" in mapping and "quantity" in mapping
            ),
            "rows_total": len(data_rows),
            "preview": data_rows[:PREVIEW_ROWS],
            "_quality": (
                header_score,
                len(data_rows),
                -sheet_order,
            ),
        }
        if best is None or candidate["_quality"] > best["_quality"]:
            best = candidate

    if best is None:
        raise ValueError("Файл не содержит данных.")

    best.pop("_quality", None)
    return best


def _parse_quantity(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    text = text.replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


VALID_DUPLICATE_POLICIES = {"sum", "max", "last", "reject"}


def duplicate_oem_stats(
    path: str | Path,
    mapping: dict[str, int],
    sheet_name: str | None = None,
) -> dict[str, int]:
    if "oem" not in mapping:
        return {
            "duplicate_oems": 0,
            "duplicate_rows": 0,
        }

    inspection = inspect_stock_file(
        path,
        sheet_name=sheet_name,
    )
    _, _, rows = read_table(
        path,
        sheet_name=inspection["sheet_name"],
    )
    oem_index = mapping["oem"]
    counts: dict[str, int] = {}

    for row in rows[inspection["header_row"]:]:
        if not any(str(x or "").strip() for x in row):
            continue
        oem = str(
            row[oem_index] if oem_index < len(row) else ""
        ).strip()
        if not oem:
            continue
        key = oem.casefold()
        counts[key] = counts.get(key, 0) + 1

    duplicate_oems = sum(
        1 for count in counts.values() if count > 1
    )
    duplicate_rows = sum(
        count - 1
        for count in counts.values()
        if count > 1
    )
    return {
        "duplicate_oems": duplicate_oems,
        "duplicate_rows": duplicate_rows,
    }


def prepare_stock_rows(
    path: str | Path,
    mapping: dict[str, int] | None = None,
    sheet_name: str | None = None,
    duplicate_oem_policy: str = "sum",
) -> dict[str, Any]:
    duplicate_oem_policy = str(
        duplicate_oem_policy or "sum"
    ).strip().lower()
    if duplicate_oem_policy not in VALID_DUPLICATE_POLICIES:
        raise ValueError(
            "Некорректная политика повторяющихся OEM."
        )

    inspection = inspect_stock_file(
        path,
        sheet_name=sheet_name,
    )
    actual_mapping = mapping or inspection["mapping"]
    if "oem" not in actual_mapping or "quantity" not in actual_mapping:
        raise ValueError("Не удалось определить колонки OEM и количества.")

    _, _, rows = read_table(
        path,
        sheet_name=inspection["sheet_name"],
    )
    data_rows = rows[inspection["header_row"]:]

    parsed: list[dict[str, Any]] = []
    errors = 0

    for row_number, row in enumerate(
        data_rows,
        start=inspection["header_row"] + 1,
    ):
        if not any(str(x or "").strip() for x in row):
            continue
        try:
            oem_index = actual_mapping["oem"]
            qty_index = actual_mapping["quantity"]
            oem = str(
                row[oem_index] if oem_index < len(row) else ""
            ).strip()
            qty_raw = (
                row[qty_index]
                if qty_index < len(row)
                else None
            )
            qty = _parse_quantity(qty_raw)
            if not oem or qty is None:
                errors += 1
                continue

            # Negative supplier stock means "nothing available"
            # for client availability. Raw source data is still
            # retained in source_record for audit.
            normalized_qty = max(float(qty), 0.0)

            manufacturer = None
            name = None
            price = None
            if "manufacturer" in actual_mapping:
                idx = actual_mapping["manufacturer"]
                raw_manufacturer = row[idx] if idx < len(row) else None
                manufacturer = (
                    str(raw_manufacturer).strip()
                    if raw_manufacturer is not None
                    else ""
                ) or None
            if "name" in actual_mapping:
                idx = actual_mapping["name"]
                raw_name = row[idx] if idx < len(row) else None
                name = (
                    str(raw_name).strip()
                    if raw_name is not None
                    else ""
                ) or None
            if "price" in actual_mapping:
                idx = actual_mapping["price"]
                raw_price = row[idx] if idx < len(row) else None
                try:
                    price = float(str(raw_price).replace(" ", "").replace(",", ".")) if raw_price not in (None, "") else None
                except (TypeError, ValueError):
                    price = None

            parsed.append({
                "row_number": row_number,
                "manufacturer": manufacturer,
                "oem": oem,
                "name": name,
                "quantity": normalized_qty,
                "price_rub": price,
                "source_record": {
                    "row": row_number,
                    "raw": [
                        str(x) if x is not None else ""
                        for x in row
                    ],
                },
            })
        except Exception:
            errors += 1

    groups: dict[str, list[dict[str, Any]]] = {}
    order: list[str] = []
    for item in parsed:
        key = str(item["oem"]).casefold()
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(item)

    duplicate_oems = sum(
        1 for group in groups.values() if len(group) > 1
    )
    duplicate_rows = sum(
        len(group) - 1
        for group in groups.values()
        if len(group) > 1
    )

    if duplicate_oem_policy == "reject" and duplicate_oems:
        raise WarehouseFileError(
            "В файле найдены повторяющиеся OEM: "
            f"{duplicate_oems} позиций / "
            f"{duplicate_rows} лишних строк."
        )

    prepared: list[dict[str, Any]] = []
    for key in order:
        group = groups[key]
        if len(group) == 1:
            prepared.append(group[0])
            continue

        if duplicate_oem_policy == "last":
            chosen = dict(group[-1])
        elif duplicate_oem_policy == "max":
            chosen = dict(
                max(
                    group,
                    key=lambda item: float(
                        item["quantity"]
                    ),
                )
            )
        else:
            chosen = dict(group[0])
            chosen["quantity"] = sum(
                float(item["quantity"])
                for item in group
            )

        chosen["manufacturer"] = next(
            (
                item.get("manufacturer")
                for item in group
                if item.get("manufacturer")
            ),
            chosen.get("manufacturer"),
        )
        chosen["name"] = next(
            (
                item.get("name")
                for item in group
                if item.get("name")
            ),
            chosen.get("name"),
        )
        chosen["source_record"] = {
            "duplicate_policy": duplicate_oem_policy,
            "aggregated_rows": [
                item["source_record"]
                for item in group
            ],
        }
        prepared.append(chosen)

    return {
        **inspection,
        "mapping": actual_mapping,
        "rows_parsed": len(parsed),
        "rows_success": len(prepared),
        "rows_error": errors,
        "duplicate_oems": duplicate_oems,
        "duplicate_rows": duplicate_rows,
        "duplicate_oem_policy": duplicate_oem_policy,
        "items": prepared,
    }
