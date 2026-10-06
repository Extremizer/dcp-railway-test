#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Railway-side OEM price-list maintenance.

Goals:
- keep price-list analysis/import next to /data/oem_reference.db;
- never import legacy prices;
- dry-run first, apply only after an explicit confirmation;
- keep immutable source SHA/provenance;
- make a consistent SQLite backup before every apply;
- emit compact JSON reports into Railway logs and /data/import_reports.

This module is intentionally isolated from customer-facing business logic.
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import shutil
import sqlite3
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from typing import Any

from fastapi import APIRouter, File, Form, Header, HTTPException, UploadFile
from fastapi.responses import HTMLResponse
from pydantic import BaseModel
from openpyxl import load_workbook

import supplier_admin_auth

log = logging.getLogger("oem_import_maintenance")

DATA_DIR = Path(os.getenv("RAILWAY_VOLUME_MOUNT_PATH", "/data"))
DB_PATH = DATA_DIR / "oem_reference.db"
SOURCE_DIR = DATA_DIR / "import_sources"
REPORT_DIR = DATA_DIR / "import_reports"
BACKUP_DIR = DATA_DIR / "backups"
JOB_DIR = DATA_DIR / "import_jobs"
MAX_UPLOAD_BYTES = 80 * 1024 * 1024
_WORKER = ThreadPoolExecutor(max_workers=1, thread_name_prefix="oem-import")
_JOB_LOCK = threading.Lock()

for _p in (SOURCE_DIR, REPORT_DIR, BACKUP_DIR, JOB_DIR):
    _p.mkdir(parents=True, exist_ok=True)

# Reusable profiles. More can be added without changing the runner.
PROFILES: dict[str, dict[str, Any]] = {
    "KAWASAKI_DEALER_2020": {
        "label": "KAWASAKI DEALER 2020",
        "manufacturer": "KAWASAKI",
        "source_year": 2020,
        "source_kind": "legacy_price_list",
        "trust_level": "legacy",
        "header_row": 1,
        "oem_header": "part_number",
        "name_header": "part_name",
        "replacement_header": "superseded #",
        "create_target_only": True,
    },
    "HONDA_DEALER_2021": {
        "label": "HONDA DEALER 2021",
        "manufacturer": "HONDA",
        "source_year": 2021,
        "source_kind": "legacy_price_list",
        "trust_level": "legacy",
        "header_row": 1,
        "oem_header": "part_number",
        "name_header": "part_name",
        "replacement_header": "superseded #",
        "create_target_only": True,
    },
    "BRP_CANAM_SEADOO_DEALER_2021": {
        "label": "CAN AM, SEA DOO DEALER 2021",
        "manufacturer": "BRP",
        "source_year": 2021,
        "source_kind": "legacy_price_list",
        "trust_level": "legacy",
        "header_row": 1,
        "oem_header": "part_number",
        "name_header": "part_name",
        "replacement_header": "ЗАМЕНА",
        "create_target_only": True,
    },
}

router = APIRouter()


class ApplyRequest(BaseModel):
    confirmation: str


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _safe_filename(name: str) -> str:
    raw = Path(str(name or "upload.xlsx")).name
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", raw).strip("._")
    return stem or "upload.xlsx"


def _normalize_oem(value: Any) -> str:
    s = str(value or "").replace("\xa0", " ").strip().upper()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    return re.sub(r"\s+", " ", s)


def _clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ").strip())


def _norm_header(value: Any) -> str:
    return _clean_text(value).casefold()


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _job_path(job_id: str) -> Path:
    return JOB_DIR / f"{job_id}.json"


def _report_path(job_id: str, mode: str) -> Path:
    return REPORT_DIR / f"{job_id}_{mode}.json"


def _write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, path)


def _load_job(job_id: str) -> dict[str, Any]:
    path = _job_path(job_id)
    if not path.exists():
        raise KeyError(job_id)
    return json.loads(path.read_text(encoding="utf-8"))


def _save_job(job: dict[str, Any]) -> None:
    with _JOB_LOCK:
        job["updated_at"] = _utc_now()
        _write_json_atomic(_job_path(job["job_id"]), job)


def _require_admin(token: str | None) -> None:
    try:
        supplier_admin_auth.require_web_admin_token(token)
    except PermissionError as exc:
        raise HTTPException(status_code=401, detail=str(exc))


def _resolve_headers(ws, profile: dict[str, Any]) -> dict[str, int]:
    row = next(ws.iter_rows(min_row=int(profile["header_row"]), max_row=int(profile["header_row"]), values_only=True))
    lookup = {_norm_header(v): i for i, v in enumerate(row) if _clean_text(v)}
    result: dict[str, int] = {}
    for logical, config_key in (
        ("oem", "oem_header"),
        ("name", "name_header"),
        ("replacement", "replacement_header"),
    ):
        wanted = _norm_header(profile.get(config_key))
        if not wanted:
            continue
        if wanted not in lookup:
            raise ValueError(f"required header not found: {profile.get(config_key)}")
        result[logical] = lookup[wanted]
    return result


def _parse_source(path: Path, profile: dict[str, Any]) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        idx = _resolve_headers(ws, profile)
        names: dict[str, set[str]] = {}
        pairs: set[tuple[str, str]] = set()
        source_rows = 0
        invalid_self: set[tuple[str, str]] = set()
        start_row = int(profile["header_row"]) + 1
        for row in ws.iter_rows(min_row=start_row, values_only=True):
            oem = _normalize_oem(row[idx["oem"]] if idx["oem"] < len(row) else None)
            if not oem:
                continue
            source_rows += 1
            name = _clean_text(row[idx["name"]] if idx.get("name", -1) < len(row) and "name" in idx else None)
            if name:
                names.setdefault(oem, set()).add(name)
            repl = _normalize_oem(row[idx["replacement"]] if idx.get("replacement", -1) < len(row) and "replacement" in idx else None)
            if repl and repl not in {"0", "N/A", "NA", "NONE", "-"}:
                if repl == oem:
                    invalid_self.add((oem, repl))
                else:
                    pairs.add((oem, repl))
        return {
            "source_rows": source_rows,
            "names": names,
            "pairs": pairs,
            "invalid_self": invalid_self,
        }
    finally:
        wb.close()


def _existing_relations(conn: sqlite3.Connection, relation_type: str) -> set[tuple[str, str]]:
    return {
        (str(r[0]), str(r[1]))
        for r in conn.execute(
            "SELECT from_oem,to_oem FROM oem_relations WHERE relation_type=?",
            (relation_type,),
        )
    }


def _build_dry_run(source_path: Path, profile_key: str) -> dict[str, Any]:
    if profile_key not in PROFILES:
        raise ValueError("unknown profile")
    profile = PROFILES[profile_key]
    parsed = _parse_source(source_path, profile)
    names: dict[str, set[str]] = parsed["names"]
    pairs: set[tuple[str, str]] = parsed["pairs"]
    source_oems = set(names)
    conflicts = {o: sorted(v) for o, v in names.items() if len(v) > 1}
    safe_names = {o: next(iter(v)) for o, v in names.items() if len(v) == 1}
    target_oems = {b for _, b in pairs}
    target_only = target_oems - source_oems

    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        prod = {
            r["oem"]: r
            for r in conn.execute(
                """SELECT oem,canonical_name,manufacturer,item_type,
                          actual_weight_kg,volume_weight_kg
                     FROM oem_reference"""
            )
        }
        existing = source_oems & set(prod)
        new_source = source_oems - set(prod)
        target_only_existing = target_only & set(prod)
        target_only_new = target_only - set(prod)
        exact = _existing_relations(conn, "replacement")
        new_edges = pairs - exact
        integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
        schema_version_row = conn.execute(
            "SELECT value FROM oem_reference_meta WHERE key='schema_version'"
        ).fetchone()
        schema_version = schema_version_row[0] if schema_version_row else None

    observations = sum(len(v) for v in names.values())
    manufacturer = str(profile.get("manufacturer") or "").strip() or None
    report = {
        "mode": "dry-run",
        "profile_key": profile_key,
        "source_key": profile_key,
        "source_name": profile["label"],
        "source_year": profile.get("source_year"),
        "source_kind": profile.get("source_kind"),
        "trust_level": profile.get("trust_level"),
        "manufacturer": manufacturer,
        "source_sha256": _sha256(source_path),
        "source_size_bytes": source_path.stat().st_size,
        "source_rows": parsed["source_rows"],
        "source_unique_oems": len(source_oems),
        "source_distinct_name_observations": observations,
        "source_oems_with_single_name": len(safe_names),
        "source_name_conflict_oems": len(conflicts),
        "source_name_conflicts": conflicts,
        "invalid_self_replacement_pairs": len(parsed["invalid_self"]),
        "valid_replacement_pairs_from_source": len(pairs),
        "replacement_pairs_already_exact": len(pairs & exact),
        "would_add_new_replacement_edges": len(new_edges),
        "replacement_target_only_oems": len(target_only),
        "replacement_target_only_already_in_production": len(target_only_existing),
        "would_insert_minimal_target_only_oems": len(target_only_new) if profile.get("create_target_only") else 0,
        "production_oems_before": len(prod),
        "already_in_production": len(existing),
        "would_insert_new_source_oems": len(new_source),
        "would_insert_total_new_oem_rows": len(new_source) + (len(target_only_new) if profile.get("create_target_only") else 0),
        "production_oems_after_if_applied": len(prod) + len(new_source) + (len(target_only_new) if profile.get("create_target_only") else 0),
        "would_add_source_name_observations": observations,
        "would_set_canonical_name_safe_total": len(safe_names),
        "would_leave_canonical_name_unset_due_conflict": len(conflicts),
        "manufacturer_observations_named_oems": len(source_oems) if manufacturer else 0,
        "manufacturer_observations_target_only": len(target_only_new) if manufacturer and profile.get("create_target_only") else 0,
        "existing_item_type_preserved": sum(1 for o in existing if (prod[o]["item_type"] or "").strip()),
        "existing_actual_weight_preserved": sum(1 for o in existing if prod[o]["actual_weight_kg"] is not None and prod[o]["actual_weight_kg"] > 0),
        "existing_volume_weight_preserved": sum(1 for o in existing if prod[o]["volume_weight_kg"] is not None and prod[o]["volume_weight_kg"] > 0),
        "weight_observations": 0,
        "cross_relations": 0,
        "analog_relations": 0,
        "prices_read_or_saved": False,
        "database_integrity_before": integrity,
        "schema_version": schema_version,
        "database_size_bytes_before": DB_PATH.stat().st_size,
        "generated_at": _utc_now(),
    }
    if integrity != "ok":
        raise RuntimeError(f"database integrity failed before dry-run: {integrity}")
    return report


def _backup_database(job_id: str) -> tuple[Path, str]:
    free = shutil.disk_usage(DATA_DIR).free
    db_size = DB_PATH.stat().st_size
    if free < db_size * 2 + 100 * 1024 * 1024:
        raise RuntimeError("insufficient free space for safe database backup/import")
    backup = BACKUP_DIR / f"oem_reference.pre_{job_id}_{int(time.time())}.db"
    src = sqlite3.connect(DB_PATH)
    dst = sqlite3.connect(backup)
    try:
        src.backup(dst)
    finally:
        dst.close()
        src.close()
    with sqlite3.connect(backup) as c:
        integrity = c.execute("PRAGMA integrity_check").fetchone()[0]
    if integrity != "ok":
        backup.unlink(missing_ok=True)
        raise RuntimeError("backup integrity check failed")
    return backup, _sha256(backup)


def _apply(job: dict[str, Any], dry: dict[str, Any]) -> dict[str, Any]:
    profile_key = job["profile_key"]
    profile = PROFILES[profile_key]
    source_path = Path(job["source_path"])
    if _sha256(source_path) != dry["source_sha256"]:
        raise RuntimeError("source file changed after dry-run")

    # Re-run exact dry-run immediately before apply so production drift cannot be hidden.
    fresh = _build_dry_run(source_path, profile_key)
    stable_keys = (
        "source_sha256",
        "source_unique_oems",
        "source_distinct_name_observations",
        "source_name_conflict_oems",
        "valid_replacement_pairs_from_source",
        "replacement_pairs_already_exact",
        "would_add_new_replacement_edges",
        "would_insert_total_new_oem_rows",
        "production_oems_before",
    )
    drift = {
        k: {"dry_run": dry.get(k), "pre_apply": fresh.get(k)}
        for k in stable_keys
        if dry.get(k) != fresh.get(k)
    }
    if drift:
        raise RuntimeError("production/source drift detected; new dry-run required: " + json.dumps(drift))

    parsed = _parse_source(source_path, profile)
    names: dict[str, set[str]] = parsed["names"]
    pairs: set[tuple[str, str]] = parsed["pairs"]
    source_oems = set(names)
    conflicts = {o for o, v in names.items() if len(v) > 1}
    safe_names = {o: next(iter(v)) for o, v in names.items() if len(v) == 1}
    target_only = {b for _, b in pairs} - source_oems
    manufacturer = str(profile.get("manufacturer") or "").strip() or None

    backup_path, backup_sha = _backup_database(job["job_id"])

    conn = sqlite3.connect(DB_PATH, timeout=60)
    conn.row_factory = sqlite3.Row
    try:
        prod = {r["oem"]: r for r in conn.execute("SELECT * FROM oem_reference")}
        existing = source_oems & set(prod)
        new_source = source_oems - set(prod)
        target_only_new = target_only - set(prod)
        target_only_existing = target_only & set(prod)
        existing_exact = _existing_relations(conn, "replacement")

        conn.execute("PRAGMA foreign_keys=ON")
        conn.execute("BEGIN IMMEDIATE")
        conn.execute(
            """INSERT INTO reference_sources
               (source_key,source_name,manufacturer,source_year,source_kind,trust_level,source_sha256,notes)
               VALUES(?,?,?,?,?,?,?,?)
               ON CONFLICT(source_key) DO UPDATE SET
                 source_name=excluded.source_name,
                 manufacturer=excluded.manufacturer,
                 source_year=excluded.source_year,
                 source_kind=excluded.source_kind,
                 trust_level=excluded.trust_level,
                 source_sha256=excluded.source_sha256""",
            (
                profile_key,
                profile["label"],
                manufacturer,
                profile.get("source_year"),
                profile.get("source_kind"),
                profile.get("trust_level"),
                dry["source_sha256"],
                "Railway maintenance import. Legacy prices are intentionally excluded.",
            ),
        )

        inserted_source = 0
        for o in sorted(new_source):
            nm = safe_names.get(o)
            conflict = o in conflicts
            cur = conn.execute(
                """INSERT OR IGNORE INTO oem_reference(
                     oem,item_type,actual_weight_kg,volume_weight_kg,weight_state,
                     direct_observation_count,linked_group_count,needs_review,
                     type_conflict,actual_weight_conflict,volume_weight_conflict,
                     canonical_name,manufacturer,name_state,manufacturer_state,
                     name_source_count,manufacturer_source_count)
                   VALUES(?,NULL,NULL,NULL,'NO_DIRECT_WEIGHT',0,0,?,0,0,0,?,?,?, ?,?,?)""",
                (
                    o,
                    1 if conflict else 0,
                    None if conflict else nm,
                    manufacturer,
                    "CONFLICT" if conflict else ("LEGACY_SINGLE_SOURCE" if nm else None),
                    "LEGACY_SINGLE_SOURCE" if manufacturer else None,
                    1 if names.get(o) else 0,
                    1 if manufacturer else 0,
                ),
            )
            inserted_source += cur.rowcount

        inserted_targets = 0
        if profile.get("create_target_only"):
            for o in sorted(target_only_new):
                cur = conn.execute(
                    """INSERT OR IGNORE INTO oem_reference(
                         oem,item_type,actual_weight_kg,volume_weight_kg,weight_state,
                         direct_observation_count,linked_group_count,needs_review,
                         type_conflict,actual_weight_conflict,volume_weight_conflict,
                         canonical_name,manufacturer,name_state,manufacturer_state,
                         name_source_count,manufacturer_source_count)
                       VALUES(?,NULL,NULL,NULL,'NO_DIRECT_WEIGHT',0,0,0,0,0,0,NULL,?,NULL,?,0,?)""",
                    (
                        o,
                        manufacturer,
                        "LEGACY_SINGLE_SOURCE" if manufacturer else None,
                        1 if manufacturer else 0,
                    ),
                )
                inserted_targets += cur.rowcount

        set_existing_names = 0
        for o in sorted(existing):
            row = prod[o]
            if o in safe_names and not (row["canonical_name"] or "").strip():
                set_existing_names += conn.execute(
                    """UPDATE oem_reference
                          SET canonical_name=?,name_state='LEGACY_SINGLE_SOURCE'
                        WHERE oem=? AND (canonical_name IS NULL OR trim(canonical_name)='')""",
                    (safe_names[o], o),
                ).rowcount
            if manufacturer and not (row["manufacturer"] or "").strip():
                conn.execute(
                    """UPDATE oem_reference
                          SET manufacturer=?,manufacturer_state='LEGACY_SINGLE_SOURCE'
                        WHERE oem=? AND (manufacturer IS NULL OR trim(manufacturer)='')""",
                    (manufacturer, o),
                )

        if manufacturer:
            for o in sorted(target_only_existing):
                row = prod[o]
                if not (row["manufacturer"] or "").strip():
                    conn.execute(
                        """UPDATE oem_reference
                              SET manufacturer=?,manufacturer_state='LEGACY_SINGLE_SOURCE'
                            WHERE oem=? AND (manufacturer IS NULL OR trim(manufacturer)='')""",
                        (manufacturer, o),
                    )

        name_obs = 0
        manufacturer_obs = 0
        for o in sorted(source_oems):
            status = "review" if o in conflicts else "active"
            for nm in sorted(names[o]):
                name_obs += conn.execute(
                    """INSERT OR IGNORE INTO oem_fact_observations(
                         oem,fact_type,value_text,source_key,trust_level,status)
                       VALUES(?,'name',?,?,?,?)""",
                    (o, nm, profile_key, profile.get("trust_level"), status),
                ).rowcount
            if manufacturer:
                manufacturer_obs += conn.execute(
                    """INSERT OR IGNORE INTO oem_fact_observations(
                         oem,fact_type,value_text,source_key,trust_level,status)
                       VALUES(?,'manufacturer',?,?,?,'active')""",
                    (o, manufacturer, profile_key, profile.get("trust_level")),
                ).rowcount

        if manufacturer and profile.get("create_target_only"):
            for o in sorted(target_only_new):
                manufacturer_obs += conn.execute(
                    """INSERT OR IGNORE INTO oem_fact_observations(
                         oem,fact_type,value_text,source_key,trust_level,status)
                       VALUES(?,'manufacturer',?,?,?,'active')""",
                    (o, manufacturer, profile_key, profile.get("trust_level")),
                ).rowcount

        rel_obs = 0
        rel_rows = 0
        for a, b in sorted(pairs):
            rel_obs += conn.execute(
                """INSERT OR IGNORE INTO oem_relation_observations(
                     from_oem,to_oem,relation_type,source_key,trust_level,status,direction_verified)
                   VALUES(?,?,'replacement',?,?,'active',1)""",
                (a, b, profile_key, profile.get("trust_level")),
            ).rowcount
            pre = 1 if (a, b) in existing_exact else 0
            state = "LEGACY_CONFIRMED_PRE_V2" if pre else "LEGACY_SINGLE_SOURCE"
            rel_rows += conn.execute(
                """INSERT OR IGNORE INTO oem_relations(
                     from_oem,to_oem,relation_type,relation_state,source_count,needs_review,pre_v2_known)
                   VALUES(?,?,'replacement',?,1,0,?)""",
                (a, b, state, pre),
            ).rowcount

        conn.execute(
            """UPDATE oem_reference
                  SET name_source_count=(
                      SELECT COUNT(DISTINCT source_key)
                        FROM oem_fact_observations f
                       WHERE f.oem=oem_reference.oem
                         AND f.fact_type='name'
                         AND f.status IN ('active','review'))
                WHERE oem IN (
                    SELECT DISTINCT oem FROM oem_fact_observations
                     WHERE source_key=? AND fact_type='name')""",
            (profile_key,),
        )
        conn.execute(
            """UPDATE oem_reference
                  SET manufacturer_source_count=(
                      SELECT COUNT(DISTINCT source_key)
                        FROM oem_fact_observations f
                       WHERE f.oem=oem_reference.oem
                         AND f.fact_type='manufacturer'
                         AND f.status='active')
                WHERE oem IN (
                    SELECT DISTINCT oem FROM oem_fact_observations
                     WHERE source_key=? AND fact_type='manufacturer')""",
            (profile_key,),
        )
        for o in conflicts:
            row = prod.get(o)
            if row is None or not (row["canonical_name"] or "").strip():
                conn.execute(
                    "UPDATE oem_reference SET needs_review=1,name_state='CONFLICT',canonical_name=NULL WHERE oem=?",
                    (o,),
                )
            else:
                conn.execute("UPDATE oem_reference SET needs_review=1 WHERE oem=?", (o,))

        conn.execute(
            """UPDATE oem_relations
                  SET source_count=(
                      SELECT COUNT(DISTINCT source_key)
                        FROM oem_relation_observations x
                       WHERE x.from_oem=oem_relations.from_oem
                         AND x.to_oem=oem_relations.to_oem
                         AND x.relation_type=oem_relations.relation_type
                         AND x.status='active')
                WHERE relation_type='replacement'"""
        )
        conn.commit()

        integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
        if integrity != "ok":
            raise RuntimeError(f"post-import integrity failed: {integrity}")

        final = {
            "mode": "apply",
            "profile_key": profile_key,
            "source_sha256": dry["source_sha256"],
            "backup_path": str(backup_path),
            "backup_sha256": backup_sha,
            "inserted_source_oems": inserted_source,
            "inserted_target_only_oems": inserted_targets,
            "inserted_total_oems": inserted_source + inserted_targets,
            "set_existing_canonical_names": set_existing_names,
            "inserted_name_observations": name_obs,
            "inserted_manufacturer_observations": manufacturer_obs,
            "inserted_relation_observations": rel_obs,
            "inserted_relation_rows": rel_rows,
            "production_oems_after": conn.execute("SELECT COUNT(*) FROM oem_reference").fetchone()[0],
            "database_integrity_after": integrity,
            "database_sha256_after": _sha256(DB_PATH),
            "database_size_bytes_after": DB_PATH.stat().st_size,
            "prices_read_or_saved": False,
            "weight_observations": 0,
            "completed_at": _utc_now(),
        }
        return final
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def _run_dry_job(job_id: str) -> None:
    try:
        job = _load_job(job_id)
        job["status"] = "dry_running"
        _save_job(job)
        report = _build_dry_run(Path(job["source_path"]), job["profile_key"])
        report["job_id"] = job_id
        _write_json_atomic(_report_path(job_id, "dryrun"), report)
        job["status"] = "dry_run_complete"
        job["dry_run_report"] = str(_report_path(job_id, "dryrun"))
        _save_job(job)
        log.warning("OEM_IMPORT_DRYRUN %s", json.dumps(report, ensure_ascii=False, separators=(",", ":")))
    except Exception as exc:
        log.exception("OEM_IMPORT_DRYRUN_FAILED job_id=%s", job_id)
        try:
            job = _load_job(job_id)
            job["status"] = "failed"
            job["error"] = repr(exc)
            _save_job(job)
        except Exception:
            pass


def _run_apply_job(job_id: str) -> None:
    try:
        job = _load_job(job_id)
        dry_path = Path(job.get("dry_run_report") or "")
        if not dry_path.exists():
            raise RuntimeError("dry-run report missing")
        dry = json.loads(dry_path.read_text(encoding="utf-8"))
        job["status"] = "applying"
        _save_job(job)
        report = _apply(job, dry)
        report["job_id"] = job_id
        _write_json_atomic(_report_path(job_id, "apply"), report)
        job["status"] = "applied"
        job["apply_report"] = str(_report_path(job_id, "apply"))
        _save_job(job)
        log.warning("OEM_IMPORT_APPLY %s", json.dumps(report, ensure_ascii=False, separators=(",", ":")))
    except Exception as exc:
        log.exception("OEM_IMPORT_APPLY_FAILED job_id=%s", job_id)
        try:
            job = _load_job(job_id)
            job["status"] = "failed"
            job["error"] = repr(exc)
            _save_job(job)
        except Exception:
            pass


def _public_job(job: dict[str, Any]) -> dict[str, Any]:
    return {
        k: v
        for k, v in job.items()
        if k not in {"source_path"}
    }


@router.get("/admin/oem-import", response_class=HTMLResponse)
def admin_page() -> str:
    options = "".join(
        f'<option value="{escape(k)}">{escape(v["label"])}</option>'
        for k, v in PROFILES.items()
    )
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width">
<title>OEM Import Maintenance</title>
<style>
body{{font-family:system-ui;background:#11131a;color:#eef2ff;margin:0}}main{{max-width:900px;margin:36px auto;padding:0 18px}}
.card{{background:#1b1f2a;border:1px solid #343a4a;border-radius:14px;padding:18px;margin:14px 0}}
input,select,button{{font:inherit;padding:10px;border-radius:8px;border:1px solid #475067;background:#10131b;color:#fff}}
button{{cursor:pointer}}button.primary{{background:#6d45d7}}pre{{white-space:pre-wrap;word-break:break-word}}
.warn{{color:#ffcf70}}.ok{{color:#6be49b}}.bad{{color:#ff7d7d}}small{{color:#aab2c5}}
</style></head><body><main>
<h1>OEM Import Maintenance</h1>
<div class="card">
<p><b>Admin token</b></p><input id="token" type="password" style="width:100%" autocomplete="off">
<p><b>Profile</b></p><select id="profile">{options}</select>
<p><b>XLSX</b></p><input id="file" type="file" accept=".xlsx">
<p><button class="primary" onclick="upload()">Upload + automatic dry-run</button></p>
<small>Legacy prices are never imported. Dry-run is mandatory before apply.</small>
</div>
<div class="card"><button onclick="refresh()">Refresh jobs</button><div id="jobs"></div></div>
<script>
const tokenInput=document.getElementById('token');
if(location.hash.startsWith('#token=')){{
  sessionStorage.setItem('oemAdminToken',decodeURIComponent(location.hash.slice(7)));
  history.replaceState(null,'',location.pathname);
}}
tokenInput.value=sessionStorage.getItem('oemAdminToken')||'';
tokenInput.onchange=()=>sessionStorage.setItem('oemAdminToken',tokenInput.value);
function H(){{return {{'X-Extremizer-Admin-Token':tokenInput.value}}}}
async function upload(){{
 const f=document.getElementById('file').files[0]; if(!f) return alert('Choose XLSX');
 const fd=new FormData(); fd.append('profile_key',document.getElementById('profile').value); fd.append('file',f);
 const r=await fetch('/api/admin/oem-import/upload',{{method:'POST',headers:H(),body:fd}});
 alert(await r.text()); refresh();
}}
async function applyJob(id){{
 if(!confirm('Apply '+id+' to production? A SQLite backup will be created first.')) return;
 const r=await fetch('/api/admin/oem-import/jobs/'+id+'/apply',{{method:'POST',headers:{{...H(),'Content-Type':'application/json'}},body:JSON.stringify({{confirmation:'APPLY '+id}})}});
 alert(await r.text()); refresh();
}}
async function refresh(){{
 const r=await fetch('/api/admin/oem-import/jobs',{{headers:H()}});
 if(!r.ok){{document.getElementById('jobs').innerHTML='<p class="bad">'+await r.text()+'</p>';return}}
 const data=await r.json(); let h='';
 for(const j of data.jobs){{
   h+='<div style="border-top:1px solid #343a4a;padding:12px 0"><b>'+j.job_id+'</b> · '+j.profile_key+' · <span>'+j.status+'</span>';
   if(j.status==='dry_run_complete') h+=' <button onclick="applyJob(\\''+j.job_id+'\\')">Apply</button>';
   if(j.error) h+='<pre class="bad">'+j.error+'</pre>';
   h+='</div>';
 }}
 document.getElementById('jobs').innerHTML=h||'<p>No jobs yet.</p>';
}}
setInterval(refresh,5000); refresh();
</script></main></body></html>"""


@router.post("/api/admin/oem-import/upload")
async def upload_source(
    profile_key: str = Form(...),
    file: UploadFile = File(...),
    x_extremizer_admin_token: str | None = Header(default=None),
) -> dict[str, Any]:
    _require_admin(x_extremizer_admin_token)
    if profile_key not in PROFILES:
        raise HTTPException(status_code=400, detail="unknown profile")
    name = _safe_filename(file.filename or "upload.xlsx")
    if not name.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="only .xlsx is accepted")
    job_id = f"{profile_key.lower()}-{uuid.uuid4().hex[:10]}"
    dest = SOURCE_DIR / f"{job_id}__{name}"
    total = 0
    with dest.open("wb") as out:
        while True:
            chunk = await file.read(1024 * 1024)
            if not chunk:
                break
            total += len(chunk)
            if total > MAX_UPLOAD_BYTES:
                out.close()
                dest.unlink(missing_ok=True)
                raise HTTPException(status_code=413, detail="file too large")
            out.write(chunk)
    job = {
        "job_id": job_id,
        "profile_key": profile_key,
        "original_filename": name,
        "source_path": str(dest),
        "source_sha256": _sha256(dest),
        "source_size_bytes": total,
        "status": "queued_dry_run",
        "created_at": _utc_now(),
    }
    _save_job(job)
    _WORKER.submit(_run_dry_job, job_id)
    return {"ok": True, "job": _public_job(job)}


@router.get("/api/admin/oem-import/jobs")
def list_jobs(
    x_extremizer_admin_token: str | None = Header(default=None),
) -> dict[str, Any]:
    _require_admin(x_extremizer_admin_token)
    jobs: list[dict[str, Any]] = []
    for p in sorted(JOB_DIR.glob("*.json"), key=lambda x: x.stat().st_mtime, reverse=True)[:50]:
        try:
            jobs.append(_public_job(json.loads(p.read_text(encoding="utf-8"))))
        except Exception:
            continue
    return {"jobs": jobs}


@router.get("/api/admin/oem-import/jobs/{job_id}")
def get_job(
    job_id: str,
    x_extremizer_admin_token: str | None = Header(default=None),
) -> dict[str, Any]:
    _require_admin(x_extremizer_admin_token)
    try:
        job = _load_job(job_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="job not found")
    payload: dict[str, Any] = {"job": _public_job(job)}
    for key, mode in (("dry_run", "dryrun"), ("apply", "apply")):
        p = _report_path(job_id, mode)
        if p.exists():
            payload[key] = json.loads(p.read_text(encoding="utf-8"))
    return payload


@router.post("/api/admin/oem-import/jobs/{job_id}/apply")
def apply_job(
    job_id: str,
    body: ApplyRequest,
    x_extremizer_admin_token: str | None = Header(default=None),
) -> dict[str, Any]:
    _require_admin(x_extremizer_admin_token)
    try:
        job = _load_job(job_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="job not found")
    if body.confirmation != f"APPLY {job_id}":
        raise HTTPException(status_code=400, detail="confirmation mismatch")
    if job.get("status") != "dry_run_complete":
        raise HTTPException(status_code=409, detail=f"job status is {job.get('status')}")
    job["status"] = "queued_apply"
    _save_job(job)
    _WORKER.submit(_run_apply_job, job_id)
    return {"ok": True, "job_id": job_id, "status": "queued_apply"}
